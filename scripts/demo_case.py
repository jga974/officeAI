"""
Script de démonstration et de validation bout-en-bout du cas d'usage :
Analyse de f1.xls et génération de rapport_ventes.docx respectant la charte de modele.docx.
"""

from pathlib import Path
import pandas as pd
import docx
from docx.shared import Pt, RGBColor
from officeai.templates.docx_styler import DocxStyler


def run_sales_analysis():
    f1_path = Path("f1.xls")
    modele_path = Path("modele.docx")
    output_path = Path("rapport_ventes.docx")

    print(f"1. Lecture des données depuis {f1_path}...")
    df = pd.read_excel(f1_path, engine="xlrd")

    # Calculs statistiques
    total_ventes = int(df["Nombre_Ventes"].sum())
    ca_total = float(df["Total_Ventes_EUR"].sum())
    prix_moyen = float(df["Prix_Unitaire_EUR"].mean())

    df_region = df.groupby("Region", as_index=False).agg(
        Ventes=("Nombre_Ventes", "sum"),
        CA_Total_EUR=("Total_Ventes_EUR", "sum")
    ).sort_values(by="CA_Total_EUR", ascending=False)

    df_vendeur = df.groupby("Vendeur", as_index=False).agg(
        Ventes=("Nombre_Ventes", "sum"),
        CA_Total_EUR=("Total_Ventes_EUR", "sum")
    ).sort_values(by="CA_Total_EUR", ascending=False)

    print("2. Clonage de la charte graphique depuis modele.docx...")
    doc = DocxStyler.clone_template_blank(modele_path)

    # Titre principal
    title_p = doc.add_paragraph("RAPPORT D'ANALYSE DES VENTES 2026", style="Title")
    for r in title_p.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(22)
        r.font.color.rgb = RGBColor(31, 73, 125)

    sub_p = doc.add_paragraph("Synthèse automatisée des performances commerciales", style="Subtitle")
    for r in sub_p.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.italic = True
        r.font.color.rgb = RGBColor(100, 100, 100)

    # Section 1 : Synthèse
    h1 = doc.add_heading("1. Synthèse Globale", level=1)
    for r in h1.runs:
        r.font.color.rgb = RGBColor(31, 73, 125)

    doc.add_paragraph(
        f"Durant la période analysée, un volume total de {total_ventes:,} ventes a été enregistré, "
        f"générant un chiffre d'affaires global de {ca_total:,.2f} € avec un prix unitaire moyen de {prix_moyen:,.2f} €."
        .replace(",", " ")
    )

    # Section 2 : Répartition par Région
    h2 = doc.add_heading("2. Répartition par Région Commerciale", level=1)
    for r in h2.runs:
        r.font.color.rgb = RGBColor(31, 73, 125)

    DocxStyler.add_dataframe_to_doc(doc, df_region, style="Table Grid", header_bg_color="1F497D")

    # Espace
    doc.add_paragraph()

    # Section 3 : Performance par Vendeur
    h3 = doc.add_heading("3. Classement des Vendeurs", level=1)
    for r in h3.runs:
        r.font.color.rgb = RGBColor(31, 73, 125)

    DocxStyler.add_dataframe_to_doc(doc, df_vendeur, style="Table Grid", header_bg_color="1F497D")

    print(f"3. Sauvegarde du rapport Word dans {output_path}...")
    doc.save(str(output_path))
    print("[OK] Rapport genere avec succes dans rapport_ventes.docx !")


if __name__ == "__main__":
    run_sales_analysis()
