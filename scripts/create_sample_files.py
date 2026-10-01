"""
Script de génération des fichiers d'essai : f1.xls, f2.xls, f3.xls et modele.docx.
"""

from pathlib import Path
import xlwt
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def create_sample_xls(filepath: Path, base_multiplier: float = 1.0):
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Ventes_2026")

    # Styles d'en-tête
    header_style = xlwt.easyxf("font: bold on, color white; pattern: pattern solid, fore_colour dark_blue; align: horiz center")
    data_style = xlwt.easyxf("align: horiz left")
    num_style = xlwt.easyxf("align: horiz right", num_format_str="#,##0")
    currency_style = xlwt.easyxf("align: horiz right", num_format_str="#,##0.00 €")

    headers = ["Date", "Region", "Vendeur", "Produit", "Nombre_Ventes", "Prix_Unitaire_EUR", "Total_Ventes_EUR"]
    for col_idx, h in enumerate(headers):
        ws.write(0, col_idx, h, header_style)

    rows = [
        ("2026-01-05", "Nord", "Marc Dubois", "Logiciel ERP", int(12 * base_multiplier), 1500.0),
        ("2026-01-08", "Sud", "Sophie Martin", "Module CRM", int(25 * base_multiplier), 800.0),
        ("2026-01-12", "Ouest", "Thomas Leroy", "Formation IA", int(8 * base_multiplier), 2200.0),
        ("2026-01-15", "Est", "Camille Petit", "Support Annuel", int(15 * base_multiplier), 450.0),
        ("2026-01-18", "Île-de-France", "Julien Roux", "Logiciel ERP", int(30 * base_multiplier), 1500.0),
        ("2026-01-22", "Nord", "Marc Dubois", "Module CRM", int(18 * base_multiplier), 800.0),
        ("2026-01-25", "Sud", "Sophie Martin", "Formation IA", int(10 * base_multiplier), 2200.0),
        ("2026-01-28", "Île-de-France", "Camille Petit", "Logiciel ERP", int(14 * base_multiplier), 1500.0),
    ]

    for row_idx, r in enumerate(rows, start=1):
        ws.write(row_idx, 0, r[0], data_style)
        ws.write(row_idx, 1, r[1], data_style)
        ws.write(row_idx, 2, r[2], data_style)
        ws.write(row_idx, 3, r[3], data_style)
        ws.write(row_idx, 4, r[4], num_style)
        ws.write(row_idx, 5, r[5], currency_style)
        total = r[4] * r[5]
        ws.write(row_idx, 6, total, currency_style)

    # Ajuster la largeur des colonnes
    for col_idx in range(len(headers)):
        ws.col(col_idx).width = 4500

    wb.save(str(filepath))
    print(f"Fichier créé : {filepath}")


def create_sample_docx_model(filepath: Path):
    doc = docx.Document()

    # Configuration des marges
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        # En-tête
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "ENTREPRISE HORIZON | DIRECTION COMMERCIALE"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in hp.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(128, 128, 128)

        # Pied de page
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Confidentiel - Usage Interne Uniquement"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in fp.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(160, 160, 160)

    # Titre principal
    title_p = doc.add_paragraph("MODÈLE DE RAPPORT DE PERFORMANCE", style="Title")
    title_p.runs[0].font.name = "Calibri"
    title_p.runs[0].font.size = Pt(24)
    title_p.runs[0].font.color.rgb = RGBColor(31, 73, 125)  # Bleu marine

    subtitle_p = doc.add_paragraph("Charte graphique standard des rapports d'analyse", style="Subtitle")
    subtitle_p.runs[0].font.name = "Calibri"
    subtitle_p.runs[0].font.size = Pt(12)
    subtitle_p.runs[0].font.italic = True
    subtitle_p.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    doc.add_paragraph()

    # Section 1
    h1 = doc.add_heading("1. Synthèse Exécutive", level=1)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(31, 73, 125)

    p1 = doc.add_paragraph(
        "Ce paragraphe illustre la typographie du texte courant. "
        "Les rapports générés doivent respecter cette structure, avec une mise en page aérée et professionnelle."
    )
    for r in p1.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(11)

    # Section 2
    h2 = doc.add_heading("2. Indicateurs Clés et Chiffres", level=1)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(31, 73, 125)

    # Tableau modèle
    table = doc.add_table(rows=3, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    table_data = [
        ["Indicateur", "Période", "Objectif"],
        ["Volume des ventes", "Mensuel", "Atteint"],
        ["Taux de conversion", "Annuel", "En progression"],
    ]

    for row_idx, row_values in enumerate(table_data):
        row = table.rows[row_idx]
        for col_idx, val in enumerate(row_values):
            cell = row.cells[col_idx]
            cell.text = val
            if row_idx == 0:
                for cp in cell.paragraphs:
                    for crun in cp.runs:
                        crun.font.bold = True
                        crun.font.color.rgb = RGBColor(31, 73, 125)

    doc.save(str(filepath))
    print(f"Modèle Word créé : {filepath}")


if __name__ == "__main__":
    base_dir = Path.cwd()
    create_sample_xls(base_dir / "f1.xls", base_multiplier=1.0)
    create_sample_xls(base_dir / "f2.xls", base_multiplier=1.5)
    create_sample_xls(base_dir / "f3.xls", base_multiplier=0.8)
    create_sample_docx_model(base_dir / "modele.docx")
