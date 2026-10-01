"""
Tests unitaires pour le cloneur et moteur de style DocxStyler.
"""

from pathlib import Path
import docx
import pandas as pd
import pytest
from officeai.templates.docx_styler import DocxStyler


@pytest.fixture
def base_model_docx(tmp_path: Path) -> Path:
    file_path = tmp_path / "modele_base.docx"
    doc = docx.Document()
    doc.add_heading("Titre Modèle", level=0)
    doc.add_paragraph("Texte d'origine qui doit être effacé.")
    doc.save(file_path)
    return file_path


def test_docx_styler_clone_blank(base_model_docx: Path, tmp_path: Path):
    out_path = tmp_path / "nouveau_rapport.docx"
    doc = DocxStyler.clone_template_blank(base_model_docx)

    # Vérifier que les paragraphes du corps ont été vidés
    assert len(doc.paragraphs) == 0

    # Ajouter nouveau contenu avec la charte
    doc.add_heading("Nouveau Rapport Mensuel", level=1)
    doc.add_paragraph("Ce rapport hérite des styles du modèle.")

    # Ajouter un tableau pandas
    df = pd.DataFrame({
        "Produit": ["A", "B"],
        "Ventes": [150, 300]
    })
    DocxStyler.add_dataframe_to_doc(doc, df)
    doc.save(out_path)

    assert out_path.exists()

    # Recharger et vérifier
    reloaded = docx.Document(out_path)
    assert len(reloaded.paragraphs) == 2
    assert reloaded.paragraphs[0].text == "Nouveau Rapport Mensuel"
    assert len(reloaded.tables) == 1
    assert reloaded.tables[0].cell(0, 0).text == "Produit"
    assert reloaded.tables[0].cell(1, 1).text == "150"
