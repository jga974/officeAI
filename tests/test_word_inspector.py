"""
Tests unitaires pour l'inspecteur Word.
"""

from pathlib import Path
import docx
import pytest
from officeai.inspectors.word_inspector import WordInspector


@pytest.fixture
def sample_docx_model(tmp_path: Path) -> Path:
    file_path = tmp_path / "modele_test.docx"
    doc = docx.Document()
    doc.add_heading("Modèle Officiel de Rapport", level=0)
    p = doc.add_paragraph("Ceci est un document modèle pour tester la charte graphique.")
    doc.add_heading("Section 1 : Statistiques", level=1)
    doc.add_paragraph("Texte d'exemple sous la section 1.")

    # Table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Indicateur"
    table.cell(0, 1).text = "Valeur"
    table.cell(1, 0).text = "Ventes"
    table.cell(1, 1).text = "100"

    doc.save(file_path)
    return file_path


def test_word_inspector(sample_docx_model: Path):
    assert WordInspector.can_inspect(sample_docx_model)
    data = WordInspector.inspect(sample_docx_model)

    assert data["file_name"] == "modele_test.docx"
    assert data["is_jinja_template"] is False

    styles = data["paragraph_styles"]
    assert any("Heading 1" in s for s in styles)

    prompt_text = WordInspector.format_for_prompt(sample_docx_model)
    assert "modele_test.docx" in prompt_text
    assert "Heading 1" in prompt_text
