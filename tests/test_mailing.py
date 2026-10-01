"""Tests du publipostage (modele Word + fichier de donnees)."""

from pathlib import Path
import pandas as pd
import pytest
import docx
from typer.testing import CliRunner

from officeai.cli import app
from officeai.converters.pdf_converter import PdfConverter
from officeai.mailing import MailingEngine, MailingError, slugify

runner = CliRunner()
needs_soffice = pytest.mark.skipif(PdfConverter.find_soffice() is None, reason="LibreOffice absent")


def make_template(path: Path, text="Bonjour {{ prenom }} {{ nom }}, rendez-vous le {{ date_rdv }}."):
    d = docx.Document()
    d.add_paragraph(text)
    d.save(str(path))


def make_xlsx(path: Path):
    pd.DataFrame({
        "Prénom": ["Éloïse", "Jean"],
        "Nom": ["Dupont & Fils", "Martin"],
        "Date RDV": [pd.Timestamp("2026-10-05"), pd.Timestamp("2026-10-06")],
        "Montant": [1200.0, 99.5],
    }).to_excel(path, index=False)


def text_of(path: Path) -> str:
    return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)


def test_slugify():
    assert slugify("Date de naissance") == "date_de_naissance"
    assert slugify("Prénom") == "prenom"
    assert slugify("2024 CA (€)") == "_2024_ca"


def test_generates_one_document_per_row(tmp_path: Path):
    make_template(tmp_path / "lettre.docx", "Bonjour {{ prenom }} {{ nom }}, le {{ date_rdv }} : {{ montant }} EUR.")
    make_xlsx(tmp_path / "contacts.xlsx")
    res = MailingEngine(tmp_path).run(Path("lettre.docx"), Path("contacts.xlsx"), name_pattern="lettre_{nom}.docx")
    assert not res.errors and len(res.generated) == 2
    first = text_of(res.output_dir / "lettre_Dupont & Fils.docx")
    assert "Éloïse" in first and "Dupont & Fils" in first
    assert "05/10/2026" in first and "1200 EUR" in first
    assert "99.5 EUR" in text_of(res.output_dir / "lettre_Martin.docx")


def test_missing_variable_aborts_before_writing(tmp_path: Path):
    make_template(tmp_path / "lettre.docx", "Cher {{ prenom }}, ville : {{ ville }}")
    make_xlsx(tmp_path / "contacts.xlsx")
    with pytest.raises(MailingError, match="ville"):
        MailingEngine(tmp_path).run(Path("lettre.docx"), Path("contacts.xlsx"))
    assert not (tmp_path / "mailing_lettre").exists()


def test_template_without_variables_is_rejected(tmp_path: Path):
    make_template(tmp_path / "vide.docx", "Texte fixe")
    make_xlsx(tmp_path / "contacts.xlsx")
    with pytest.raises(MailingError, match="aucune variable"):
        MailingEngine(tmp_path).run(Path("vide.docx"), Path("contacts.xlsx"))


def test_never_overwrites_without_force_and_is_atomic(tmp_path: Path):
    make_template(tmp_path / "lettre.docx", "{{ prenom }}")
    make_xlsx(tmp_path / "contacts.xlsx")
    eng = MailingEngine(tmp_path)
    eng.run(Path("lettre.docx"), Path("contacts.xlsx"), name_pattern="{prenom}.docx")
    (tmp_path / "mailing_lettre" / "Jean.docx").write_bytes(b"marqueur")
    with pytest.raises(MailingError, match="existent deja"):
        eng.run(Path("lettre.docx"), Path("contacts.xlsx"), name_pattern="{prenom}.docx")
    assert (tmp_path / "mailing_lettre" / "Jean.docx").read_bytes() == b"marqueur"
    res = eng.run(Path("lettre.docx"), Path("contacts.xlsx"), name_pattern="{prenom}.docx", overwrite=True)
    assert len(res.overwritten) == 2


def test_forbidden_filename_chars_are_replaced(tmp_path: Path):
    make_template(tmp_path / "t.docx", "{{ nom }}")
    pd.DataFrame({"Nom": ['A/B:C*"D']}).to_csv(tmp_path / "d.csv", index=False)
    res = MailingEngine(tmp_path).run(Path("t.docx"), Path("d.csv"), name_pattern="{nom}.docx")
    assert [p.name for p in res.generated] == ["A_B_C__D.docx"]


def test_duplicate_names_get_suffix(tmp_path: Path):
    make_template(tmp_path / "t.docx", "{{ nom }}")
    pd.DataFrame({"Nom": ["Durand", "Durand", "Durand"]}).to_csv(tmp_path / "d.csv", index=False)
    res = MailingEngine(tmp_path).run(Path("t.docx"), Path("d.csv"), name_pattern="{nom}.docx")
    assert sorted(p.name for p in res.generated) == ["Durand.docx", "Durand_2.docx", "Durand_3.docx"]


def test_csv_semicolon_cp1252_and_empty_rows(tmp_path: Path):
    make_template(tmp_path / "t.docx", "{{ prenom }} {{ ville }}")
    (tmp_path / "d.csv").write_bytes("prenom;ville\nÉmile;Orléans\n;\nZoé;\n".encode("cp1252"))
    res = MailingEngine(tmp_path).run(Path("t.docx"), Path("d.csv"))
    assert len(res.generated) == 2
    assert text_of(res.generated[0]).strip() == "Émile Orléans"


def test_files_must_stay_in_working_dir(tmp_path: Path):
    work = tmp_path / "work"
    work.mkdir()
    make_template(tmp_path / "t.docx")
    make_xlsx(work / "d.xlsx")
    with pytest.raises(MailingError, match="hors du repertoire"):
        MailingEngine(work).run(tmp_path / "t.docx", Path("d.xlsx"))
    make_template(work / "t.docx")
    with pytest.raises(MailingError, match="hors du repertoire"):
        MailingEngine(work).run(Path("t.docx"), Path("d.xlsx"), output_dir=Path("../sortie"))


def test_cli_dry_run_then_real_run(tmp_path: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    make_template(tmp_path / "lettre.docx", "{{ prenom }} {{ nom }}")
    make_xlsx(tmp_path / "contacts.xlsx")
    r = runner.invoke(app, ["mailing", "lettre.docx", "contacts.xlsx", "--dry-run", "-n", "l_{prenom}.docx"])
    assert r.exit_code == 0 and "Aucun fichier cree" in r.output and not (tmp_path / "mailing_lettre").exists()
    r = runner.invoke(app, ["mailing", "lettre.docx", "contacts.xlsx", "-n", "l_{prenom}.docx"])
    assert r.exit_code == 0 and (tmp_path / "mailing_lettre" / "l_Jean.docx").exists()
    r = runner.invoke(app, ["mailing", "lettre.docx", "contacts.xlsx", "-n", "l_{prenom}.docx"])
    assert r.exit_code == 1 and "existent deja" in r.output


@needs_soffice
def test_pdf_option(tmp_path: Path):
    make_template(tmp_path / "t.docx", "{{ prenom }}")
    make_xlsx(tmp_path / "d.xlsx")
    res = MailingEngine(tmp_path).run(Path("t.docx"), Path("d.xlsx"), name_pattern="{prenom}.docx", to_pdf=True)
    assert len(res.pdfs) == 2 and all(p.read_bytes().startswith(b"%PDF") for p in res.pdfs)
