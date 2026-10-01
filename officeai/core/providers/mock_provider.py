"""
Fournisseur de Démonstration / Hors-ligne (Simulation & Onboarding) pour OfficeAI.
Permet la découverte et la prise en main sans nécessiter de clé d'API.
"""

from __future__ import annotations
from datetime import datetime
from pathlib import Path
import re
from officeai.core.providers.base import LLMProvider


class MockProvider(LLMProvider):
    """
    Fournisseur autonome conçu pour la démonstration, les tests et l'onboarding.
    Gère intelligemment les salutations, l'inventaire des fichiers, les métadonnées et l'analyse bureautique.
    """

    def __init__(self, api_key: str | None = None, model_name: str | None = None):
        super().__init__(api_key="demo-key", model_name="demo-offline")

    def generate_plan_and_code(
        self,
        system_instruction: str,
        user_prompt: str,
        context: str = "",
        temperature: float = 0.2,
    ) -> tuple[str, str]:
        prompt_lower = user_prompt.lower().strip()
        cwd = Path.cwd()

        # 1. Salutations simples
        if any(prompt_lower.startswith(g) or g in prompt_lower for g in ("bonjour", "salut", "hello", "coucou", "hey")):
            return (
                "Bonjour ! Je suis OfficeAI, votre assistant bureautique propulse par l'IA.\n\n"
                "Je peux vous aider a :\n"
                "1. Connaitre vos fichiers : 'liste de mes fichiers ?', 'nombre de lignes de f1.xls ?'\n"
                "2. Repondre a des questions sur vos donnees : 'qui a fait le plus de ventes dans f1.xls ?'\n"
                "3. Generer des rapports Word (.docx) avec graphiques respectant votre charte (modele.docx)\n"
                "4. Traiter des lots de fichiers (Batch) en une seule instruction.\n\n"
                "Que souhaitez-vous faire aujourd'hui ?",
                "",
            )

        # 2. Présentation et aide
        if any(h in prompt_lower for h in ("qui es-tu", "qui est tu", "aide", "help", "que sais-tu faire", "que peux-tu faire")):
            return (
                "Je suis OfficeAI, un assistant bureautique modulaire pour Windows.\n"
                "Je combine la puissance des modeles de langage (Gemini, Albert Souverain, Mistral) "
                "avec des bibliotheques Python pour analyser et produire des documents professionnels.\n\n"
                "Astuce : Utilisez la touche [TAB] pour completer les noms de vos fichiers !",
                "",
            )

        # 3. Question sur l'heure ou la date
        if any(d in prompt_lower for d in ("quelle date", "quelle heure", "quel jour", "heure")):
            now_str = datetime.now().strftime("%d/%m/%Y et il est %H:%M:%S")
            return (f"Nous sommes le {now_str}.", "")

        # 4. Demande de liste de fichiers (Inventaire du répertoire)
        if any(fl in prompt_lower for fl in ("liste de mes fichiers", "liste des fichiers", "quels fichiers", "mes fichiers", "quels sont les fichiers", "affiche les fichiers", "fichiers du dossier")):
            files = [f for f in sorted(list(cwd.glob("*"))) if f.is_file() and not f.name.startswith(".")]
            if not files:
                return ("Aucun fichier bureautique detecte dans le repertoire courant.", "")
            lines = ["Voici les fichiers presents dans votre repertoire de travail :\n"]
            for f in files:
                size_kb = round(f.stat().st_size / 1024, 1)
                lines.append(f"- **{f.name}** ({size_kb} Ko)")
            return ("\n".join(lines), "")

        local_files = [f for f in cwd.glob("*") if f.is_file() and not f.name.startswith(".")]
        named = next((f for f in sorted(local_files) if f.name.lower() in prompt_lower), None)

        # 4 a. Nombre de fichiers du repertoire
        if any(k in prompt_lower for k in ("nombre de fichier", "combien de fichier")):
            n_excel = sum(1 for f in local_files if f.suffix.lower() in (".xls", ".xlsx", ".xlsm", ".ods"))
            return (f"Le repertoire contient **{len(local_files)} fichier(s)** dont {n_excel} tableur(s).", "")

        # 4 b. Creation d'un fichier contenant une liste de fichiers (ex: "liste des fichiers excel")
        if any(w in prompt_lower for w in ("cree", "creer", "genere", "generer", "ecris", "ecrire")) and \
                "liste" in prompt_lower and "fichier" in prompt_lower and "rapport" not in prompt_lower:
            if "word" in prompt_lower:
                exts, label = (".docx", ".odt"), "Word"
            elif "excel" in prompt_lower or "exel" in prompt_lower or "tableur" in prompt_lower:
                exts, label = (".xls", ".xlsx", ".xlsm", ".ods"), "Excel"
            else:
                exts, label = None, "tous"
            m = re.search(r"[\w\-]+\.(?:txt|csv|md)", prompt_lower)
            out_name = m.group(0) if m else f"liste_fichiers_{label.lower()}.txt"
            code = (
                "from pathlib import Path\n"
                f"exts = {exts!r}\n"
                "files = sorted(p.name for p in Path('.').iterdir()\n"
                "               if p.is_file() and not p.name.startswith('.') and (exts is None or p.suffix.lower() in exts))\n"
                f"out = Path({out_name!r})\n"
                "if out.exists():\n"
                "    out = out.with_name(out.stem + '_nouveau' + out.suffix)\n"
                "out.write_text('\\n'.join(files) + '\\n', encoding='utf-8')\n"
                "print(f'{len(files)} fichier(s) ecrits dans {out}')\n"
            )
            return (f"Creation d'un fichier texte listant les fichiers ({label}) du dossier courant, sans ecraser de fichier existant.", code)

        # 4 c. Comptage par categorie dans un fichier XML nomme
        if named and named.suffix.lower() == ".xml" and any(
                w in prompt_lower for w in ("nombre", "combien", "compte", "calcul", "par type", "par ")):
            from officeai.inspectors.xml_inspector import XmlInspector
            try:
                info = XmlInspector.inspect(named)
            except Exception as exc:
                return (f"Impossible de lire {named.name} : {exc}", "")
            rec = info["record_tag"]
            attrs = list(info["elements"][rec]["attributes"]) if rec else []
            words = set(re.findall(r"\w+", prompt_lower))
            key = next((a for a in attrs if a.lower() in words), attrs[0] if attrs else None)
            if not rec or not key:
                return (f"Je n'ai pas trouve d'element repete avec attribut dans {named.name}.", "")
            code = (
                "import xml.etree.ElementTree as ET\n"
                "from collections import Counter\n"
                f"root = ET.parse({named.name!r}).getroot()\n"
                f"counts = Counter(el.get({key!r}, '(vide)') for el in root.iter({rec!r}))\n"
                f"print('Nombre de <{rec}> par {key} :')\n"
                "for k, n in counts.most_common():\n"
                "    print(f'- {k} : {n}')\n"
                "print(f'Total : {sum(counts.values())}')\n"
            )
            return (f"Lecture de {named.name} et comptage des <{rec}> par attribut '{key}'.", code)

        # 4 bis. Suppression explicite d'un fichier nomme (le script sera soumis a confirmation)
        if named and any(w in prompt_lower for w in ("supprime", "supprimer", "efface", "effacer", "delete")):
            code = (
                "from pathlib import Path\n"
                f"target = Path({named.name!r})\n"
                "target.unlink()\n"
                "print(f'Fichier supprime : {target}')\n"
            )
            return (f"Suppression du fichier {named.name} du repertoire de travail.", code)

        # 4 ter. Proprietes d'un fichier nomme : taille, date
        if named and any(w in prompt_lower for w in ("taille", "poids", "octets", "modifie", "date de")):
            st = named.stat()
            from datetime import datetime as _dt
            return (
                f"**{named.name}** : {round(st.st_size / 1024, 1)} Ko ({st.st_size} octets), "
                f"derniere modification le {_dt.fromtimestamp(st.st_mtime):%d/%m/%Y a %H:%M}.",
                "",
            )

        # 4 quater. Colonnes / feuilles d'un tableur nomme
        if named and any(w in prompt_lower for w in ("colonne", "feuille", "en-tete", "entete")):
            from officeai.inspectors.excel_inspector import ExcelInspector
            if ExcelInspector.can_inspect(named):
                data = ExcelInspector.inspect(named)
                parts = []
                for sname, sinfo in data["sheets"].items():
                    cols = ", ".join(c["name"] for c in sinfo["columns"])
                    parts.append(f"- feuille **{sname}** : {cols}")
                return (f"Colonnes de **{named.name}** :\n" + "\n".join(parts), "")
            return (f"**{named.name}** n'est pas un tableur : pas de colonnes a lister.", "")

        # 5. Demande de comptage de lignes ou colonnes d'un fichier
        if any(rc in prompt_lower for rc in ("nombre de ligne", "combien de ligne", "nombre de colonnes", "combien de colonnes")):
            # Recherche du fichier mentionné
            target_f = None
            for f in cwd.glob("*"):
                if f.is_file() and f.name.lower() in prompt_lower:
                    target_f = f
                    break
            if not target_f:
                if "f1" in prompt_lower:
                    target_f = cwd / "f1.xls"
                elif "ventes_nationales" in prompt_lower or ".ods" in prompt_lower:
                    target_f = cwd / "ventes_nationales.ods"

            if target_f and target_f.exists():
                from officeai.inspectors.excel_inspector import ExcelInspector
                if ExcelInspector.can_inspect(target_f):
                    data = ExcelInspector.inspect(target_f)
                    summaries = []
                    for sname, sinfo in data["sheets"].items():
                        summaries.append(f"feuille '{sname}' : **{sinfo['total_rows_inspected']} lignes** et **{sinfo['columns_count']} colonnes**")
                    return (f"Le fichier **{target_f.name}** comporte {', '.join(summaries)}.", "")
                else:
                    return (f"Le fichier **{target_f.name}** existe mais ce n'est pas un tableur inspectable.", "")
            else:
                return ("Veuillez preciser le nom du fichier dont vous souhaitez connaitre le nombre de lignes (ex: 'nombre de lignes de f1.xls ?').", "")

        # 6. Questions générales de culture, finance, météo
        cwd_files = [f.name.lower() for f in cwd.glob("*") if f.is_file() and not f.name.startswith(".")]
        mentions_local_file = any(fn in prompt_lower for fn in cwd_files) or any(ext in prompt_lower for ext in (".xls", ".xlsx", ".ods", ".docx", ".csv", ".pdf", ".xml"))
        is_file_operation = any(act in prompt_lower for act in ("analyse", "analyser", "traite", "traiter", "calcule", "calculer", "rapport", "document", "fichier", "dossier", "tableau", "tableur", "colonne", "lot", "batch", "ventes"))

        if not (mentions_local_file or is_file_operation):
            if "yen" in prompt_lower or any(w in prompt_lower for w in ("temperature", "meteo", "temps")):
                return (
                    "Je n'ai pas acces a internet : je ne peux pas vous donner d'information en temps reel "
                    "(meteo, cours de change). Consultez une source en ligne pour une valeur a jour.",
                    "",
                )
            else:
                return (
                    f"Reponse de l'IA a votre question : '{user_prompt}'\n\n"
                    "En tant qu'assistant IA, je reponds a toutes vos questions d'ordre general, "
                    "et j'execute localement vos commandes d'automatisation des que vous mentionnez vos fichiers locaux.",
                    "",
                )

        # 7. Question analytique simple sur les données (sans création de document Word)
        wants_document = any(w in prompt_lower for w in ("rapport", "word", "docx", "document", "cree", "generer", "enregistre", "mets", "produit"))
        is_best_seller_query = any(q in prompt_lower for q in ("plus de vente", "meilleur vendeur", "top", "qui a fait", "combien de vente"))

        source_file = "f1.xls"
        if ".ods" in prompt_lower or "ventes_nationales" in prompt_lower:
            source_file = "ventes_nationales.ods"
        elif "f2" in prompt_lower:
            source_file = "f2.xls"
        elif "f3" in prompt_lower:
            source_file = "f3.xls"

        if is_best_seller_query and not wants_document:
            explanation = (
                f"Analyse des donnees du fichier {source_file} pour identifier "
                "le collaborateur ayant realise le plus de ventes."
            )
            code = f"""import pandas as pd
from pathlib import Path

source = Path('{source_file}')
ext = source.suffix.lower()
if ext == '.ods':
    df = pd.read_excel(source, engine='odf')
elif ext == '.xls':
    df = pd.read_excel(source, engine='xlrd')
else:
    df = pd.read_excel(source, engine='openpyxl')

group_col = 'Vendeur' if 'Vendeur' in df.columns else ('Responsable' if 'Responsable' in df.columns else df.columns[1])
val_col = 'Nombre_Ventes' if 'Nombre_Ventes' in df.columns else df.columns[4]

totals = df.groupby(group_col)[val_col].sum()
top_name = totals.idxmax()
top_score = totals.max()

print(f"Le meilleur resultat a ete realise par {{top_name}} avec {{top_score}} ventes.")
"""
            return explanation, code

        if not any(w in prompt_lower for w in ("rapport", "genere", "generer", "analyse", "analyser", "traite", "traiter", "word", "docx", "cree", "creer", "synthese")):
            return (
                "Je n'ai pas bien compris votre demande. Exemples : 'liste de mes fichiers', "
                "'colonnes de f1.xls', 'taille de modele.docx', 'qui a fait le plus de ventes dans f1.xls ?', "
                "'analyse f1.xls et genere rapport_ventes.docx'.",
                "",
            )

        # 8. Demande explicite de création de rapport Word (Scénario de génération documentaire)
        explanation = (
            f"[Mode Demonstration / Hors-ligne]\n"
            f"1. Lecture des donnees dans {source_file}.\n"
            f"2. Calcul des agregats (volume des ventes, CA total, ventilation).\n"
            f"3. Clonage de la charte graphique de modele.docx (titres, polices bleues, marges).\n"
            f"4. Insertion de tableaux et sauvegarde du rapport final."
        )

        target_file = "rapport_ventes.docx"
        m = re.search(r"rapport_[\w\-]+\.docx", prompt_lower)
        if m:
            target_file = m.group(0)

        code = f"""import sys
from pathlib import Path
import pandas as pd
import docx
from docx.shared import Pt, RGBColor
from officeai.templates.docx_styler import DocxStyler

source_path = Path('{source_file}')
model_path = Path('modele.docx')
output_path = Path('{target_file}')

print(f"Chargement des donnees depuis {{source_path}}...")
ext = source_path.suffix.lower()
if ext == '.ods':
    df = pd.read_excel(source_path, engine='odf')
elif ext == '.xls':
    df = pd.read_excel(source_path, engine='xlrd')
else:
    df = pd.read_excel(source_path, engine='openpyxl')

total_ventes = int(df['Nombre_Ventes'].sum()) if 'Nombre_Ventes' in df.columns else len(df)
ca_total = float(df['Total_Ventes_EUR'].sum()) if 'Total_Ventes_EUR' in df.columns else 0.0

print(f"Calculs termines : {{total_ventes}} ventes, CA: {{ca_total:,.2f}} EUR")

group_col = 'Region' if 'Region' in df.columns else ('Direction_Regionale' if 'Direction_Regionale' in df.columns else df.columns[1])
val_col = 'Total_Ventes_EUR' if 'Total_Ventes_EUR' in df.columns else df.columns[-1]

df_synth = df.groupby(group_col, as_index=False).agg(
    Ventes=('Nombre_Ventes', 'sum'),
    CA_EUR=(val_col, 'sum')
).sort_values(by='CA_EUR', ascending=False)

if model_path.exists():
    doc = DocxStyler.clone_template_blank(model_path)
else:
    doc = docx.Document()

p_title = doc.add_paragraph("RAPPORT D'ANALYSE D'ACTIVITE", style="Title")
for r in p_title.runs:
    r.font.name = "Calibri"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(31, 73, 125)

p_sub = doc.add_paragraph("Synthese automatisee generee par OfficeAI", style="Subtitle")
for r in p_sub.runs:
    r.font.name = "Calibri"
    r.font.italic = True
    r.font.color.rgb = RGBColor(100, 100, 100)

h1 = doc.add_heading("1. Synthese Globale", level=1)
for r in h1.runs:
    r.font.color.rgb = RGBColor(31, 73, 125)

doc.add_paragraph(
    f"Durant la periode consideree, un volume de {{total_ventes:,}} ventes a ete enregistre, "
    f"pour un montant total de {{ca_total:,.2f}} EUR."
)

h2 = doc.add_heading("2. Ventilation par Secteur", level=1)
for r in h2.runs:
    r.font.color.rgb = RGBColor(31, 73, 125)

DocxStyler.add_dataframe_to_doc(doc, df_synth, style="Table Grid", header_bg_color="1F497D")

doc.save(str(output_path))
print(f"Rapport genere avec succes dans {{output_path}} !")
"""
        return explanation, code

    def auto_repair_code(
        self,
        failed_code: str,
        error_message: str,
        context: str = "",
        temperature: float = 0.1,
    ) -> str:
        return failed_code
