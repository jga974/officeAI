"""
Moteur de clonage et d'application de charte graphique Word pour OfficeAI.
Fournit des utilitaires réutilisables par les scripts générés.
"""

from __future__ import annotations
from pathlib import Path
from typing import Any
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docxtpl import DocxTemplate
import pandas as pd


class DocxStyler:
    """Outils pour cloner la charte graphique d'un document Word ou instancier un template Jinja2."""

    @staticmethod
    def render_jinja(template_path: str | Path, context: dict[str, Any], output_path: str | Path) -> Path:
        """Remplit un template Word avec docxtpl et sauvegarde le résultat."""
        t_path = Path(template_path)
        out_path = Path(output_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        tpl = DocxTemplate(t_path)
        tpl.render(context)
        tpl.save(out_path)
        return out_path

    @staticmethod
    def clone_template_blank(template_path: str | Path, output_path: str | Path | None = None) -> docx.Document:
        """
        Ouvre un document Word modèle, conserve sa mise en page, ses marges,
        ses styles, en-têtes et pieds de page, mais vide le corps du texte
        pour permettre l'écriture d'un nouveau document conforme à la charte.
        """
        t_path = Path(template_path)
        doc = docx.Document(t_path)

        # Nettoyage des paragraphes du corps du document tout en gardant la section/styles
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)

        # Nettoyage des tableaux du corps du document
        for t in list(doc.tables):
            t._element.getparent().remove(t._element)

        if output_path:
            out_path = Path(output_path)
            out_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(out_path)

        return doc

    @staticmethod
    def add_dataframe_to_doc(
        doc: docx.Document,
        df: pd.DataFrame,
        style: str | None = "Table Grid",
        header_bg_color: str | None = "1F497D",  # Bleu corporate par défaut
    ) -> docx.table.Table:
        """
        Insère un tableau élégant dans le document Word à partir d'un DataFrame pandas.
        """
        table = doc.add_table(rows=len(df) + 1, cols=len(df.columns))
        if style:
            try:
                table.style = style
            except Exception:
                pass

        # Remplissage des en-têtes
        hdr_cells = table.rows[0].cells
        for col_idx, col_name in enumerate(df.columns):
            hdr_cells[col_idx].text = str(col_name)
            # Mise en gras de l'en-tête
            for paragraph in hdr_cells[col_idx].paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
            # Couleur d'arrière-plan de l'en-tête si spécifiée
            if header_bg_color:
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{header_bg_color}"/>')
                hdr_cells[col_idx]._tc.get_or_add_tcPr().append(shading)

        # Remplissage des données
        for row_idx, row_data in enumerate(df.itertuples(index=False)):
            row_cells = table.rows[row_idx + 1].cells
            for col_idx, val in enumerate(row_data):
                if pd.isna(val):
                    val_str = ""
                elif isinstance(val, (int, float)):
                    val_str = f"{val:,.2f}".replace(",", " ").replace(".", ",") if isinstance(val, float) else str(val)
                else:
                    val_str = str(val)
                row_cells[col_idx].text = val_str

        return table
