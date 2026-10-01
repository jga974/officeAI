"""
Inspecteur de documents Word (.docx) pour extraction de charte graphique et détection de modèles.
"""

from __future__ import annotations
from pathlib import Path
from typing import Any
import docx
from docx.enum.style import WD_STYLE_TYPE
from docxtpl import DocxTemplate


class WordInspector:
    """Analyse un document Word (.docx) pour en extraire sa charte graphique et ses variables."""

    @classmethod
    def can_inspect(cls, file_path: str | Path) -> bool:
        path = Path(file_path)
        return path.suffix.lower() == ".docx"

    @classmethod
    def inspect(cls, file_path: str | Path) -> dict[str, Any]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Fichier modèle introuvable : {path}")

        doc = docx.Document(path)

        # 1. Détection des variables Jinja2 avec docxtpl
        tpl = DocxTemplate(path)
        try:
            jinja_vars = sorted(list(tpl.get_undeclared_template_variables()))
        except Exception:
            jinja_vars = []

        is_jinja_template = len(jinja_vars) > 0

        # 2. Extraction des styles de paragraphe disponibles
        paragraph_styles = []
        for s in doc.styles:
            if s.type == WD_STYLE_TYPE.PARAGRAPH:
                font_info = ""
                if s.font.name:
                    font_info += f", Police: {s.font.name}"
                if s.font.size:
                    font_info += f", Taille: {s.font.size.pt}pt"
                paragraph_styles.append(f"{s.name}{font_info}")

        # 3. Extraction des styles de tableaux
        table_styles = [s.name for s in doc.styles if s.type == WD_STYLE_TYPE.TABLE]

        # 4. Marges et dimensions de page de la première section
        page_setup = {}
        if doc.sections:
            section = doc.sections[0]
            page_setup = {
                "top_margin_cm": round(section.top_margin.cm, 2) if section.top_margin else None,
                "bottom_margin_cm": round(section.bottom_margin.cm, 2) if section.bottom_margin else None,
                "left_margin_cm": round(section.left_margin.cm, 2) if section.left_margin else None,
                "right_margin_cm": round(section.right_margin.cm, 2) if section.right_margin else None,
                "orientation": "Landscape" if section.orientation == docx.enum.section.WD_ORIENT.LANDSCAPE else "Portrait",
                "has_header": bool(section.header and section.header.paragraphs and any(p.text.strip() for p in section.header.paragraphs)),
                "has_footer": bool(section.footer and section.footer.paragraphs and any(p.text.strip() for p in section.footer.paragraphs)),
            }

        # 5. Aperçu des premiers paragraphes (structure existante)
        preview_paragraphs = []
        for p in doc.paragraphs[:8]:
            if p.text.strip():
                preview_paragraphs.append(f"[{p.style.name}] {p.text[:80]}")

        return {
            "file_name": path.name,
            "file_path": str(path.resolve()),
            "is_jinja_template": is_jinja_template,
            "jinja_variables": jinja_vars,
            "paragraph_styles": paragraph_styles,
            "table_styles": table_styles,
            "page_setup": page_setup,
            "preview_paragraphs": preview_paragraphs,
        }

    @classmethod
    def format_for_prompt(cls, file_path: str | Path) -> str:
        """
        Formate les informations de charte graphique et de template pour le prompt de Gemini.
        """
        info = cls.inspect(file_path)
        lines = [
            f"--- FICHIER WORD MODÈLE : {info['file_name']} ---",
            f"Chemin absolu : {info['file_path']}",
        ]

        if info["is_jinja_template"]:
            lines.append("Type de modèle : Modèle balisé Jinja2 (docxtpl).")
            lines.append(f"Variables Jinja2 détectées : {', '.join(info['jinja_variables'])}")
            lines.append("Recommandation : Utiliser `docxtpl.DocxTemplate` et lui passer un dictionnaire de contexte.")
        else:
            lines.append("Type de modèle : Document de référence de style / charte graphique.")
            lines.append("Recommandation : Utiliser le helper `DocxStyler.create_from_template()` pour cloner la charte graphique.")

        lines.append("\nStyles de paragraphes utilisables dans ce modèle :")
        for s in info["paragraph_styles"][:15]:
            lines.append(f"  * {s}")

        lines.append("\nStyles de tableaux utilisables :")
        for ts in info["table_styles"][:8]:
            lines.append(f"  * {ts}")

        lines.append(f"\nMise en page : Marges (G: {info['page_setup'].get('left_margin_cm')}cm, D: {info['page_setup'].get('right_margin_cm')}cm), Orientation: {info['page_setup'].get('orientation')}")
        lines.append(f"En-tête présent : {info['page_setup'].get('has_header')}, Pied de page présent : {info['page_setup'].get('has_footer')}")

        if info["preview_paragraphs"]:
            lines.append("\nExemples de paragraphes existants dans le modèle :")
            for pp in info["preview_paragraphs"]:
                lines.append(f"  * {pp}")

        return "\n".join(lines)
